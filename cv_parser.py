"""
cv_parser.py
Modul untuk mengekstrak teks mentah dari file CV berformat PDF atau DOCX.
Ini adalah tahap pertama pipeline: file -> raw text.
"""

from pathlib import Path
import pdfplumber
import docx


def extract_text_from_pdf(file_path: str) -> str:
    """Ekstrak seluruh teks dari file PDF menggunakan pdfplumber."""
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_path: str) -> str:
    """Ekstrak seluruh teks dari file DOCX menggunakan python-docx."""
    document = docx.Document(file_path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Ikutkan juga teks di dalam tabel (banyak CV pakai layout tabel)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs).strip()


def parse_cv(file_path: str) -> str:
    """
    Fungsi utama: deteksi ekstensi file lalu panggil extractor yang sesuai.
    Mengembalikan raw text. Melempar ValueError jika format tidak didukung
    atau file kosong/tidak bisa dibaca (misal hasil scan tanpa OCR).
    """
    suffix = Path(file_path).suffix.lower()

    if suffix == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif suffix == ".docx":
        text = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Format file tidak didukung: {suffix}. Gunakan .pdf atau .docx")

    if not text or len(text) < 20:
        raise ValueError(
            f"Tidak berhasil mengekstrak teks dari {file_path}. "
            "Kemungkinan file adalah hasil scan/gambar dan memerlukan OCR "
            "(di luar scope POC ini, lihat catatan 'Future Work' pada slide)."
        )

    return text
